import os
import openpyxl
from openpyxl.styles import Alignment

class UtilsExcel:
    def __init__(self, template_path, ):
        self.template_path = template_path

    def convert_to_number(self, value):
        if value is None:
            return value

        if isinstance(value, (int, float)):
            return value

        if isinstance(value, str):
            v = value.strip().replace(",", ".")

            try:
                num = float(v)
                if num.is_integer():
                    return int(num)
                return num
            except:
                return value

        return value

    def add_gush_helka_to_excel(self, data, start_row, length):

        wb = openpyxl.load_workbook(self.template_path)
        ws = wb.active

        center_align = Alignment(horizontal="center", vertical="center")

        gush_value = data["gush"]
        helka_value = data["helka"]

        gush_value = self.convert_to_number(gush_value)
        helka_value = self.convert_to_number(helka_value)

        for i in range(length):
            row = start_row + i

            cell_gush = ws.cell(row=row, column=1, value=gush_value)
            cell_gush.alignment = center_align
            cell_gush.number_format = "@"
            ws.column_dimensions[cell_gush.column_letter].width = 8

            cell_helka = ws.cell(row=row, column=2, value=helka_value)
            cell_helka.alignment = center_align
            cell_helka.number_format = "@"

        return wb

    def add_row_to_excel(self, wb, data_list, column, start_row=1, prefix=""):

        ws = wb.active

        center_align = Alignment(horizontal="center", vertical="center")

        for index, value in enumerate(data_list):
            current_row = start_row + index

            value = self.convert_to_number(f"{prefix}" + value)

            if isinstance(column, str):
                cell = ws[f"{column}{current_row}"]
                cell.value = value
                cell.alignment = center_align
                cell.number_format = "@"

            else:
                cell = ws.cell(row=current_row, column=column, value=value)
                cell.alignment = center_align
                cell.number_format = "@"

        return wb

    def add_section_data_to_excel(self, wb, section_data, start_row, start_col):

        ws = wb.active
        center_align = Alignment(horizontal="center", vertical="center")

        for item in section_data:
            if any(char.isdigit() for char in str(item.get("name_2"))):
                print ("digits found at name 2 , ignoring it ")
                item["name_2"] = ""
                item["id_2"] = ""




            if item.get("amount") == 2:
                fields = ["name_1", "id_1", "amount", "name_2", "id_2", "amount"]

            else:
                fields = ["name_1", "id_1", "amount", "name_2", "id_2"]

            col = start_col

            for field in fields:
                value = item.get(field, "")
                if field == "amount":
                    value = f"1/{value}"
                value = self.convert_to_number(value)

                cell = ws.cell(row=start_row, column=col, value=value)
                cell.alignment = center_align

                col += 1

            start_row += 1

        return wb

    def save_excel(self, wb, folder_name):
        os.makedirs(os.path.dirname(folder_name), exist_ok=True)
        wb.save(folder_name)
